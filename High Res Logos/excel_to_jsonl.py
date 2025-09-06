#!/usr/bin/env python3
"""
Excel → JSONL converter for Smooth Migration services

- Reads the workbook "Smooth Migration - Master Services .xlsx"
- Writes one JSONL per country folder inside High Res Logos/* (e.g., Canada/, USA/, UK/)
- Content format matches the WP importer expectations (Overview:/Why we recommend:/How it helps:/Link:)

Usage examples:

  python3 excel_to_jsonl.py \
    --excel "High Res Logos/Smooth Migration - Master Services .xlsx" \
    --append --verbose

  python3 excel_to_jsonl.py --sheets Canada,USA --overwrite

Requirements:
  pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple, Optional

try:
    from openpyxl import load_workbook
except Exception as exc:  # pragma: no cover
    sys.stderr.write(
        "[ERROR] Missing dependency: openpyxl. Install with: pip install openpyxl\n"
    )
    raise


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert Smooth Migration Excel workbook into per-country JSONL files"
    )
    default_excel = (
        Path(__file__).parent / "Smooth Migration - Master Services .xlsx"
    )

    parser.add_argument(
        "--excel",
        type=Path,
        default=default_excel,
        help="Path to the Excel workbook (.xlsx)",
    )
    parser.add_argument(
        "--base",
        type=Path,
        default=Path(__file__).parent,
        help="Base folder containing country subfolders (default: script directory)",
    )
    parser.add_argument(
        "--sheets",
        type=str,
        default="",
        help="Comma-separated sheet names to include (defaults to all)",
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--append",
        action="store_true",
        help="Append to existing *_context.jsonl files (default if file exists)",
    )
    mode.add_argument(
        "--overwrite",
        action="store_true",
        help="Overwrite existing *_context.jsonl files",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Analyze and print actions without writing files",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="Print detailed processing logs",
    )
    return parser.parse_args()


def verbose_print(enabled: bool, message: str) -> None:
    if enabled:
        print(message)


def normalize_country_name(name: str) -> str:
    text = name.strip()
    text_lower = text.lower()
    if text_lower in {"usa", "us", "u.s.a", "u.s.", "united states", "united states of america"}:
        return "USA"
    if text_lower in {"uk", "u.k.", "united kingdom", "great britain", "britain"}:
        return "UK"
    if text_lower in {"south africa", "rsa"}:
        return "South Africa"
    if text_lower == "canada":
        return "Canada"
    if text_lower == "australia":
        return "Australia"
    # Fallback to title-case for unknown but present sheet names
    return text.title()


def country_to_jsonl_filename(country: str) -> str:
    # Prefer long-form names for UK/USA in the filename to match existing files
    mapping = {
        "UK": "united-kingdom_context.jsonl",
        "USA": "united-states_context.jsonl",
        "South Africa": "south_africa_context.jsonl",  # matches existing underscore style
        "Canada": "canada_context.jsonl",
        "Australia": "australia_context.jsonl",
    }
    slug = mapping.get(country)
    if slug:
        return slug
    # Generic hyphenated fallback
    return f"{re.sub(r'[^a-z0-9]+', '-', country.lower()).strip('-')}_context.jsonl"


def find_or_target_jsonl(country_dir: Path, country: str, verbose: bool) -> Path:
    target = country_dir / country_to_jsonl_filename(country)
    matches = list(country_dir.glob("*_context.jsonl"))
    if target.exists():
        verbose_print(verbose, f"Using existing file: {target}")
        return target
    if matches:
        # Prefer matching our target name; otherwise, first existing
        verbose_print(verbose, f"Found existing context file(s): {', '.join(str(m) for m in matches)}")
        return target if target in matches else matches[0]
    return target


def normalize_header(text: str) -> str:
    t = (text or "").strip().lower()
    # Strip trailing punctuation/colons and collapse spaces
    t = re.sub(r"[:\u2013\u2014\-\s]+$", "", t)  # remove trailing colon/dash
    t = re.sub(r"\s+", " ", t)
    replacements = {
        "country": "country",
        "partner": "partner",
        "service": "partner",
        "brand": "partner",
        "company": "partner",
        "category": "category",
        "subcategory": "subcategory",
        "overview": "overview",
        "why we recommend": "why",
        "why do we recommend": "why",
        "why recommend": "why",
        "how it helps": "how",
        "how does it help": "how",
        "how this helps": "how",
        "link": "link",
        "url": "link",
        "widget": "widget",
        "embed": "widget",
    }
    return replacements.get(t, t)

HEADER_CANON_KEYS = {
    "country",
    "partner",
    "category",
    "subcategory",
    "overview",
    "why",
    "how",
    "link",
    "widget",
}


def detect_header_row(sheet) -> Tuple[int, Dict[int, str]]:
    """Scan the first 25 rows to find the likeliest header row by recognized columns."""
    best_idx = -1
    best_score = -1
    best_mapping: Dict[int, str] = {}

    for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if r_idx > 24:
            break
        if not row:
            continue
        mapping: Dict[int, str] = {}
        score = 0
        for idx, cell in enumerate(row):
            if cell is None:
                continue
            key = normalize_header(str(cell))
            if key in HEADER_CANON_KEYS:
                score += 1
                mapping[idx] = key
        # Must include at least 'partner' or 'overview' to be a valid header
        if "partner" not in mapping.values() and "overview" not in mapping.values():
            continue
        if score > best_score:
            best_score = score
            best_idx = r_idx
            best_mapping = mapping

    # If none found, fallback to first non-empty row with loose mapping
    if best_idx == -1:
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if any(cell not in (None, "") for cell in row or []):
                mapping: Dict[int, str] = {}
                for idx, cell in enumerate(row or []):
                    key = normalize_header(str(cell or ""))
                    if key:
                        mapping[idx] = key
                return r_idx, mapping
    return best_idx, best_mapping


def row_to_record(row: Tuple, headers: Dict[int, str]) -> Dict[str, str]:
    record: Dict[str, str] = {}
    for idx, key in headers.items():
        val = row[idx] if idx < len(row) else None
        if val is None:
            continue
        text = str(val).strip()
        if text == "":
            continue
        record[key] = text
    return record


def compose_text(overview: str, why: str, how: str, link: Optional[str]) -> str:
    parts: List[str] = []
    if overview:
        parts.append(f"Overview: {overview}")
    if why:
        parts.append(f"Why we recommend: {why}")
    if how:
        parts.append(f"How it helps: {how}")
    if link:
        parts.append(f"Link: {link}")
    return "\n".join(parts)


def record_to_jsonl_line(
    country: str,
    rec: Dict[str, str],
    source_path: Path,
) -> str:
    partner = rec.get("partner", "").strip()
    category = rec.get("category", "").strip()
    overview = rec.get("overview", "").strip()
    why = rec.get("why", "").strip()
    how = rec.get("how", "").strip()
    link = rec.get("link", "").strip() or None
    widget = rec.get("widget", "").strip() or None

    text = compose_text(overview, why, how, link)

    payload = {
        "text": text,
        "metadata": {
            "country": country,
            "partner": partner,
            "category": category,
            "subcategory": rec.get("subcategory", "").strip() or None,
            "link": link,
            "widget": widget,
            "ingested_at": datetime.utcnow().isoformat() + "Z",
            "source": f"{source_path.name}::{country}",
        },
    }
    # Remove nulls from metadata
    payload["metadata"] = {k: v for k, v in payload["metadata"].items() if v}
    return json.dumps(payload, ensure_ascii=False)


def process_sheet(
    sheet,
    country_dir: Path,
    country_name_from_sheet: str,
    out_path: Path,
    overwrite: bool,
    dry_run: bool,
    verbose: bool,
    source_excel: Path,
) -> Tuple[int, int]:
    header_row_idx, headers = detect_header_row(sheet)
    if not headers or header_row_idx < 0:
        verbose_print(verbose, f"  Skipping '{sheet.title}': no suitable header row detected")
        return (0, 0)

    # Verbose: show detected header mapping summary
    header_keys = ", ".join(sorted(set(headers.values())))
    verbose_print(verbose, f"  Detected header row {header_row_idx + 1} with keys: {header_keys}")

    written = 0
    skipped = 0

    lines: List[str] = []
    for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if r_idx <= header_row_idx:
            continue
        rec = row_to_record(row, headers)
        if not rec:
            continue

        # Prefer sheet-implied country over cell content
        country = normalize_country_name(country_name_from_sheet)
        partner = rec.get("partner", "").strip()
        if not partner:
            # Try to infer from alternative keys that might have slipped through
            for alt_key in ("brand", "company", "service"):
                if alt_key in rec and rec[alt_key].strip():
                    rec["partner"] = rec[alt_key].strip()
                    partner = rec["partner"]
                    break
        if not partner:
            skipped += 1
            continue

        lines.append(record_to_jsonl_line(country, rec, source_excel))
        written += 1

    if not lines:
        verbose_print(verbose, f"  No valid rows in '{sheet.title}'")
        return (0, skipped)

    # IO
    if dry_run:
        print(f"[DRY-RUN] Would write {written} JSONL record(s) to: {out_path}")
        return (written, skipped)

    if out_path.exists() and overwrite:
        verbose_print(verbose, f"  Overwriting existing file: {out_path}")
        out_path.write_text("", encoding="utf-8")

    # Ensure directory exists
    out_path.parent.mkdir(parents=True, exist_ok=True)

    mode = "a" if out_path.exists() and not overwrite else "w"
    with out_path.open(mode, encoding="utf-8") as fh:
        for line in lines:
            fh.write(line)
            fh.write("\n")

    return (written, skipped)


def main() -> int:
    args = parse_args()

    base_dir = args.base.resolve()
    excel_path = args.excel.resolve()
    if not excel_path.exists():
        sys.stderr.write(f"[ERROR] Excel file not found: {excel_path}\n")
        return 2

    wb = load_workbook(filename=str(excel_path), data_only=True)

    include_sheets: Optional[List[str]] = None
    if args.sheets.strip():
        include_sheets = [s.strip() for s in args.sheets.split(",") if s.strip()]

    total_written = 0
    total_skipped = 0

    for sheet_name in wb.sheetnames:
        if include_sheets and sheet_name not in include_sheets:
            continue

        # Normalize country and locate the directory
        country = normalize_country_name(sheet_name)
        country_dir = base_dir / country
        if not country_dir.is_dir():
            verbose_print(
                args.verbose,
                f"Skipping sheet '{sheet_name}': country folder not found at {country_dir}",
            )
            continue

        out_path = find_or_target_jsonl(country_dir, country, args.verbose)
        sheet = wb[sheet_name]
        print(f"Processing sheet: {sheet_name} → {out_path.relative_to(base_dir)}")

        written, skipped = process_sheet(
            sheet=sheet,
            country_dir=country_dir,
            country_name_from_sheet=sheet_name,
            out_path=out_path,
            overwrite=args.overwrite and not args.append,
            dry_run=args.dry_run,
            verbose=args.verbose,
            source_excel=excel_path,
        )

        total_written += written
        total_skipped += skipped

    print(
        f"Done. Records written: {total_written}. Rows skipped (no partner): {total_skipped}."
    )
    if args.dry_run:
        print("No files were modified due to --dry-run.")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())


